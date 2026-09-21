from dataclasses import dataclass, field

import openpyxl


@dataclass(frozen=True, slots=True)
class Cell:
    text: str | None = None


@dataclass(frozen=True, slots=True)
class TableScheme:
    schema: dict[tuple[int, int], Cell]
    merges: list[tuple[int, int, int, int]] = field(default_factory=list[tuple[int, int, int, int]])
    row_sizes: list[tuple[int, int]] = field(default_factory=list[tuple[int, int]])


@dataclass(frozen=True, slots=True)
class Schemes:
    REQUEST = TableScheme(
        schema={
            (0, 10): Cell("Характеристики"),
            (0, 13): Cell("Доп. информация"),
            (1, 0): Cell("Номенклатура 1С"),
            (1, 1): Cell("П1"),
            (1, 2): Cell("Р1"),
            (1, 3): Cell("П2"),
            (1, 4): Cell("Р2"),
            (1, 5): Cell("П3"),
            (1, 6): Cell("Р3"),
            (1, 7): Cell("П4"),
            (1, 8): Cell("Герметик"),
            (1, 9): Cell("Тип изделия"),
            (1, 10): Cell("X"),
            (1, 11): Cell("Y"),
            (1, 12): Cell("Кол-во"),
            (1, 13): Cell("Маркировка"),
            (1, 14): Cell("ПЗ"),
            (1, 15): Cell("ШК"),
            (1, 16): Cell("Номер фигуры"),
            (1, 17): Cell("Уточнения")
        },
        merges=[
            (0, 0, 0, 9),
            (0, 10, 0, 12),
            (0, 13, 0, 17)
        ],
        row_sizes=[
            (0, 15),
            (1, 25)
        ]
    )
    CALC = TableScheme(
        schema={
            (0, 0): Cell("ID"),
            (0, 1): Cell("Кол-во"),
            (0, 2): Cell("материалы"),
            (0, 7): Cell("Рамка от края"),
            (0, 8): Cell("Ед. Изм."),
            (0, 9): Cell("размеры"),
            (0, 11): Cell("Припуск на обработку"),
            (0, 16): Cell("Вращ."),
            (0, 17): Cell("Приор."),
            (0, 18): Cell("Предп."),
            (0, 19): Cell("Пирам."),
            (0, 20): Cell("Пир1Р"),
            (0, 21): Cell("Пир2М"),
            (0, 22): Cell("Пир2Р"),
            (0, 23): Cell("Пир2М"),
            (0, 24): Cell("Этикетки"),
            (0, 25): Cell("Заказ"),
            (0, 26): Cell("Заказчик"),
            (0, 27): Cell("Дата раскроя"),
            (0, 28): Cell("Текст рамки"),
            (0, 29): Cell("Отв. В рамке"),
            (0, 30): Cell("Имя фигуры"),
            (0, 32): Cell("Shape Parameters"),
            (0, 42): Cell("Shape Trims"),
            (0, 46): Cell("Shape\nElaboration"),
            (0, 47): Cell("Piece Notes"),
            (1, 2): Cell("1материал"),
            (1, 3): Cell("1рамка"),
            (1, 4): Cell("2материал"),
            (1, 5): Cell("2рамка"),
            (1, 6): Cell("3материал"),
            (1, 9): Cell("X"),
            (1, 10): Cell("Y"),
            (1, 11): Cell("Общий"),
            (1, 12): Cell("X1"),
            (1, 13): Cell("Y1"),
            (1, 14): Cell("X2"),
            (1, 15): Cell("Y2"),
            (1, 30): Cell("Импорт"),
            (1, 31): Cell("Сохран."),
            (1, 32): Cell("Par1"),
            (1, 33): Cell("Par2"),
            (1, 34): Cell("Par3"),
            (1, 35): Cell("Par4"),
            (1, 36): Cell("Par5"),
            (1, 37): Cell("Par6"),
            (1, 38): Cell("Par 7"),
            (1, 39): Cell("Par 8"),
            (1, 40): Cell("Par 9"),
            (1, 41): Cell("Par 10"),
            (1, 42): Cell("Rif X1"),
            (1, 43): Cell("Rif Y1"),
            (1, 44): Cell("Rif X2"),
            (1, 45): Cell("Rif Y2"),
            (1, 47): Cell("Note"),
            (1, 48): Cell("Note 1"),
            (1, 49): Cell("Note 2"),
            (1, 50): Cell("Note 3"),
            (1, 51): Cell("Note 4"),
            (1, 52): Cell("Note 5"),
            (1, 53): Cell("Note 6"),
            (1, 54): Cell("Note 7"),
            (1, 55): Cell("Note 8"),
            (1, 56): Cell("Note 9"),
            (1, 57): Cell("Note 10"),
            (1, 58): Cell("Note 11"),
            (1, 59): Cell("Note 12"),
            (1, 60): Cell("Note 13"),
            (1, 61): Cell("Note 14"),
            (1, 62): Cell("Note 15"),
            (1, 63): Cell("Note 16"),
            (1, 64): Cell("Note 17"),
            (1, 65): Cell("Note 18"),
            (1, 66): Cell("Note 19"),
            (1, 67): Cell("Note 20"),
        },
        merges=[
            (0, 42, 0, 45),
            (0, 47, 0, 67),
            (0, 11, 0, 15),
            (0, 2, 0, 6),
            (0, 9, 0, 10),
            (0, 30, 0, 31),
            (0, 32, 0, 41),
        ]
    )


class SchemeMaker:
    def __init__(self, scheme: TableScheme) -> None:
        self.scheme = scheme
        self.wb = openpyxl.Workbook()

        default_sheet = self.wb.active
        self.wb.remove(default_sheet)

        self.active_sheet = self.wb.create_sheet(title="Sheet1")

    def make_scheme(self) -> None:
        for (row_idx, col_idx), cell in self.scheme.schema.items():
            cell = self.active_sheet.cell(row=row_idx + 1, column=col_idx + 1, value=cell.text)

        for m_range in self.scheme.merges:
            minr, minc, maxr, maxc = m_range
            self.active_sheet.merge_cells(
                start_row=minr + 1,
                start_column=minc + 1,
                end_row=maxr + 1,
                end_column=maxc + 1,
            )

        for idx, height in self.scheme.row_sizes:
            self.active_sheet.row_dimensions[idx + 1].height = height
