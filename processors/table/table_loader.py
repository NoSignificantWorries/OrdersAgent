import re
from collections.abc import Generator
from dataclasses import dataclass
from enum import Enum
from io import BytesIO
from pathlib import Path

import openpyxl
import xlrd

NUMBER = re.compile(r"^\s*((\d+)([.,]\d+)?)\s*$")
WHITESPACE = re.compile(r"\s*")


def number_to_int(number: str) -> int | None:
    match = NUMBER.match(number)
    if match:
        number = match.group(2)
        return int(number)
    return None


def clean_str(text: str | None) -> str | None:
    if text is None:
        return None
    value = re.sub(WHITESPACE, "", text)
    if value == "":
        return None
    return value


def parse_value(value) -> int | str | None:
    if value is None:
        return None
    if isinstance(value, int | float):
        return int(value)
    if isinstance(value, str):
        clean_value = clean_str(value)
        if clean_value is None:
            return None
        num = number_to_int(clean_value)
        if num is None:
            return value
        return num
    return str(value)


class TableType(str, Enum):
    XLS = ".xls"
    XLSX = ".xlsx"


@dataclass(slots=True)
class WorkbookResults:
    name: str | None
    fmt: TableType
    with_metadata: bool
    workbook: xlrd.Book | openpyxl.Workbook


@dataclass(frozen=True, slots=True)
class Cell:
    value: int | str | None
    row: int
    col: int
    merged: bool = False
    parent: tuple[int, int] | None = None


@dataclass(frozen=True, slots=True)
class Sheet:
    name: str
    nrows: int
    ncols: int
    cells: Generator[Cell, None, None]


@dataclass(slots=True)
class Workbook:
    name: str | None
    fmt: TableType
    with_metadata: bool
    sheets: Generator[Sheet, None, None]


class TableLoader:
    @staticmethod
    def get_name(src: BytesIO | Path) -> str | None:
        if isinstance(src, BytesIO):
            return getattr(src, "name", None)
        elif isinstance(src, Path):
            return src.name

    @staticmethod
    def _open_by_type(src: BytesIO | Path, fmt: TableType, with_metadata: bool = True) -> xlrd.Book | openpyxl.Workbook:
        match fmt:
            case TableType.XLS:
                if isinstance(src, BytesIO):
                    return xlrd.open_workbook(file_contents=src.getvalue(), formatting_info=with_metadata)
                else:
                    return xlrd.open_workbook(str(src), formatting_info=with_metadata)
            case TableType.XLSX:
                if isinstance(src, BytesIO):
                    src.seek(0)
                return openpyxl.load_workbook(src, data_only=not with_metadata)
            case _:
                raise ValueError(f"Unsupported format: {fmt}")

    @staticmethod
    def _open_with_metadata_check(src: BytesIO | Path, fmt: TableType) -> WorkbookResults:
        for with_metadata in (True, False):
            try:
                wb = TableLoader._open_by_type(src, fmt, with_metadata)
                return WorkbookResults(TableLoader.get_name(src), fmt, with_metadata, wb)
            except Exception as err:
                if not with_metadata:
                    raise
                # TODO: logging with WARN for metadata errors

    @staticmethod
    def _detect_type(src: BytesIO | Path) -> TableType:
        if isinstance(src, BytesIO):
            src.seek(0)
            header = src.read(8)
            src.seek(0)
            if header.startswith(b"PK\x03\x04"):
                return TableType.XLSX
            if header.startswith(b"\xd0\xcf\x11\xe0"):
                return TableType.XLS
            raise ValueError("Unknown binary format: not .xls or .xlsx")
        elif isinstance(src, Path):
            filetype = src.suffix.lower()
            if filetype == ".xls":
                return TableType.XLS
            elif filetype == ".xlsx":
                return TableType.XLSX
            else:
                raise ValueError(f"Unsupported file type '*{filetype}'! Only *.xls and *.xlsx supported.")
        else:
            raise TypeError(f"Unsupported data type '{type(src).__name__}'! Use Path or BytesIO objects instead.")

    @staticmethod
    def _iter_xls_merge_ranges(sheet) -> Generator[tuple[int, int, int, int], None, None]:
        yield from sheet.merged_cells

    @staticmethod
    def _iter_xlsx_merge_ranges(sheet) -> Generator[tuple[int, int, int, int], None, None]:
        for rng in sheet.merged_cells.ranges:
            yield rng.min_row - 1, rng.max_row, rng.min_col - 1, rng.max_col

    @staticmethod
    def _build_merged_map(ranges: Generator[tuple[int, int, int, int]]) -> dict[tuple[int, int], tuple[int, int]]:
        merged: dict[tuple[int, int], tuple[int, int]] = {}
        for rlo, rhi, clo, chi in ranges:
            parent = (rlo, clo)
            for r in range(rlo, rhi):
                for c in range(clo, chi):
                    if (r, c) != parent:
                        merged[(r, c)] = parent
        return merged

    @staticmethod
    def _make_cell(value, row: int, col: int, parent: tuple[int, int] | None = None) -> Cell | None:
        value = parse_value(value)
        if parent is not None:
            return Cell(value=None, row=row, col=col, merged=True, parent=parent)

        if value is None or value == "":
            return None

        return Cell(value=value, row=row, col=col)


    @staticmethod
    def _iter_xls_sheets(wb: xlrd.Book) -> Generator[Sheet, None, None]:
        for sheetname in wb.sheet_names():
            sheet = wb[sheetname]
            merged = TableLoader._build_merged_map(TableLoader._iter_xls_merge_ranges(sheet))
            yield Sheet(
                name=sheetname,
                nrows=sheet.nrows,
                ncols=sheet.ncols,
                cells=TableLoader._iter_xls_cells(sheet, merged)
            )

    @staticmethod
    def _iter_xls_cells(sheet, merged: dict[tuple[int, int], tuple[int, int]]) -> Generator[Cell, None, None]:
        for row in range(sheet.nrows):
            for col in range(sheet.ncols):
                value = sheet.cell_value(row, col)
                cell = TableLoader._make_cell(value, row, col, merged.get((row, col), None))
                if cell is not None:
                    yield cell

    @staticmethod
    def _iter_xlsx_sheets(wb: openpyxl.Workbook) -> Generator[Sheet, None, None]:
        for sheetname in wb.sheetnames:
            sheet = wb[sheetname]
            merged = TableLoader._build_merged_map(TableLoader._iter_xlsx_merge_ranges(sheet))
            yield Sheet(
                name=sheetname,
                nrows=sheet.max_row,
                ncols=sheet.max_column,
                cells=TableLoader._iter_xlsx_cells(sheet, merged)
            )

    @staticmethod
    def _iter_xlsx_cells(sheet, merged: dict[tuple[int, int], tuple[int, int]]) -> Generator[Cell, None, None]:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                cell_object = TableLoader._make_cell(
                    value=value,
                    row=cell.row - 1,
                    col=cell.column - 1,
                    parent=merged.get((cell.row -1, cell.column - 1), None)
                )
                if cell_object is not None:
                    yield cell_object

    @staticmethod
    def load(src: BytesIO | Path) -> Workbook:
        fmt = TableLoader._detect_type(src)
        workbook = TableLoader._open_with_metadata_check(src, fmt)

        match fmt:
            case TableType.XLS:
                sheets = TableLoader._iter_xls_sheets(workbook.workbook)
            case TableType.XLSX:
                sheets = TableLoader._iter_xlsx_sheets(workbook.workbook)
            case _:
                raise ValueError(f"Unsupported format: {fmt}")

        return Workbook(
            name=workbook.name,
            fmt=fmt,
            with_metadata=workbook.with_metadata,
            sheets=sheets
        )
