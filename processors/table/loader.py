import re
from collections.abc import Generator
from dataclasses import dataclass
from enum import Enum
from io import BytesIO
from pathlib import Path

import openpyxl
import xlrd


@dataclass(slots=True)
class Cell:
    row: int
    col: int
    value: int | str | tuple[int, int] | None = None


class SparseTable:
    def __init__(self, nrows: int, ncols: int) -> None:
        self.nrows = nrows
        self.ncols = ncols

        self.cells: list[list[Cell]] = [[Cell(r, c) for c in range(self.ncols)] for r in range(self.nrows)]

    @property
    def empty(self) -> bool:
        return self.nrows == 0 or self.ncols == 0

    def add_cell(self, cell: Cell) -> None:
        self.cells[cell.row][cell.col] = cell

    def from_cells_generator(self, generator: Generator[Cell, None, None]) -> None:
        for cell in generator:
            self.add_cell(cell)

    def normalize(self) -> tuple[list[int], list[int]]:
        keep_rows: set[int] | list[int] = set()
        keep_cols: set[int] | list[int] = set()
        for r in range(self.nrows):
            for c in range(self.ncols):
                if self.cells[r][c].value is not None:
                    keep_cols.add(c)
                    keep_rows.add(r)
        keep_rows = sorted(keep_rows)
        keep_cols = sorted(keep_cols)
        self.cells = [[self.cells[r][c] for c in keep_cols] for r in keep_rows]
        self.nrows = len(keep_rows)
        self.ncols = len(keep_cols)

        for i in range(self.nrows):
            for j in range(self.ncols):
                self.cells[i][j].row = i
                self.cells[i][j].col = j

        return keep_rows, keep_cols

    def get_cell(self, row: int, col: int) -> Cell:
        return self.cells[row][col]

    def iter_cells(self) -> Generator[tuple[int, int, Cell], None, None]:
        for i in range(self.nrows):
            for j in range(self.ncols):
                cell = self.get_cell(i, j)
                yield i, j, cell


NUMBER = re.compile(r"^\s*((\d+)([.,]\d+)?)\s*$")
WHITESPACE = re.compile(r"\s*")


def number_to_int(number: str) -> int | None:
    match = NUMBER.match(number)
    if match:
        number = match.group(2)
        return int(number)
    return None


def clean_str(text: str | None) -> str | None:
    if text is None:
        return None
    value = re.sub(WHITESPACE, "", text)
    if value == "":
        return None
    return value


def parse_value(value) -> int | str | None:
    if value is None:
        return None
    if isinstance(value, int | float):
        return int(value)
    if isinstance(value, str):
        clean_value = clean_str(value)
        if clean_value is None:
            return None
        num = number_to_int(clean_value)
        if num is None:
            return value
        return num
    return str(value)


class TableType(str, Enum):
    XLS = ".xls"
    XLSX = ".xlsx"


@dataclass(slots=True)
class WorkbookResults:
    name: str | None
    fmt: TableType
    with_metadata: bool
    workbook: xlrd.Book | openpyxl.Workbook


@dataclass(frozen=True, slots=True)
class Sheet:
    name: str
    nrows: int
    ncols: int
    table: SparseTable


@dataclass(slots=True)
class Workbook:
    name: str | None
    fmt: TableType
    sheets:list[Sheet]


class TableLoader:
    @staticmethod
    def _get_name(src: BytesIO | Path) -> str | None:
        if isinstance(src, BytesIO):
            return getattr(src, "name", None)
        elif isinstance(src, Path):
            return src.name

    @staticmethod
    def _open_by_type(src: BytesIO | Path, fmt: TableType) -> xlrd.Book | openpyxl.Workbook:
        match fmt:
            case TableType.XLS:
                if isinstance(src, BytesIO):
                    return xlrd.open_workbook(file_contents=src.getvalue(), on_demand=True, formatting_info=False)
                else:
                    return xlrd.open_workbook(str(src), on_demand=True, formatting_info=False)
            case TableType.XLSX:
                if isinstance(src, BytesIO):
                    src.seek(0)
                return openpyxl.load_workbook(src, read_only=True, data_only=True)
            case _:
                raise ValueError(f"Unsupported format: {fmt}")

    @staticmethod
    def _close_by_type(wb: xlrd.Book | openpyxl.Workbook) -> None:
        if isinstance(wb, xlrd.Book):
            wb.release_resources()
        elif isinstance(wb, openpyxl.Workbook):
            wb.close()
        else:
            raise TypeError(f"Unsupported object '{type(wb).__name__}'")

    # @staticmethod
    # def  close_by_type(wb: xlrd.Book | openpyxl.Workbook, fmt: TableType) -> None:
    #     match fmt:
    #         case TableType.XLS:
    #             wb.close()

    @staticmethod
    def _detect_type(src: BytesIO | Path) -> TableType:
        if isinstance(src, BytesIO):
            src.seek(0)
            header = src.read(8)
            src.seek(0)
            if header.startswith(b"PK\x03\x04"):
                return TableType.XLSX
            if header.startswith(b"\xd0\xcf\x11\xe0"):
                return TableType.XLS
            raise ValueError("Unknown binary format: not .xls or .xlsx")
        elif isinstance(src, Path):
            filetype = src.suffix.lower()
            if filetype == ".xls":
                return TableType.XLS
            elif filetype == ".xlsx":
                return TableType.XLSX
            else:
                raise ValueError(f"Unsupported file type '*{filetype}'! Only *.xls and *.xlsx supported.")
        else:
            raise TypeError(f"Unsupported data type '{type(src).__name__}'! Use Path or BytesIO objects instead.")

    @staticmethod
    def _make_cell(value, row: int, col: int) -> Cell | None:
        value = parse_value(value)

        if value is None or value == "":
            return None

        return Cell(value=value, row=row, col=col)


    @staticmethod
    def _make_xls_sheets(wb: xlrd.Book) -> list[Sheet]:
        sheets: list[Sheet] = []
        for sheetname in wb.sheet_names():
            sheet = wb[sheetname]
            sparse_table = SparseTable(sheet.nrows, sheet.ncols)
            sparse_table.from_cells_generator(TableLoader._iter_xls_cells(sheet))
            new_sheet = Sheet(
                name=sheetname,
                nrows=sheet.nrows,
                ncols=sheet.ncols,
                table=sparse_table
            )
            sheets.append(new_sheet)
            wb.unload_sheet(sheetname)
        return sheets

    @staticmethod
    def _iter_xls_cells(sheet) -> Generator[Cell, None, None]:
        for row in range(sheet.nrows):
            for col in range(sheet.ncols):
                value = sheet.cell_value(row, col)
                cell = TableLoader._make_cell(value, row, col)
                if cell is not None:
                    yield cell

    @staticmethod
    def _make_xlsx_sheets(wb: openpyxl.Workbook) -> list[Sheet]:
        sheets: list[Sheet] = []
        for sheetname in wb.sheetnames:
            sheet = wb[sheetname]
            sparse_table = SparseTable(sheet.max_row, sheet.max_column)
            sparse_table.from_cells_generator(TableLoader._iter_xlsx_cells(sheet))
            new_sheet = Sheet(
                name=sheetname,
                nrows=sheet.max_row,
                ncols=sheet.max_column,
                table=sparse_table
            )
            sheets.append(new_sheet)
        return sheets

    @staticmethod
    def _iter_xlsx_cells(sheet) -> Generator[Cell, None, None]:
        for row_idx, row in enumerate(sheet.iter_rows()):
            for col_idx, cell in enumerate(row):
                value = cell.value
                cell_object = TableLoader._make_cell(
                    value=value,
                    row=row_idx,
                    col=col_idx,
                )
                if cell_object is not None:
                    yield cell_object

    @staticmethod
    def load(src: BytesIO | Path) -> Workbook:
        fmt = TableLoader._detect_type(src)
        workbook = TableLoader._open_by_type(src, fmt)

        match fmt:
            case TableType.XLS:
                sheets = TableLoader._make_xls_sheets(workbook)
            case TableType.XLSX:
                sheets = TableLoader._make_xlsx_sheets(workbook)
            case _:
                raise ValueError(f"Unsupported format: {fmt}")

        TableLoader._close_by_type(workbook)

        return Workbook(
            name=TableLoader._get_name(src),
            fmt=fmt,
            sheets=sheets
        )
