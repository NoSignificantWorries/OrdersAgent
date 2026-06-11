import re
from pathlib import Path

import openpyxl
import xlrd

from materials import DELIMETERS, ParserV2
from table import TableWorker, make_callculation_xlsx, make_request_xlsx, tp2


def test_callculation_table():
    input = Path("../private/tables/1249A.xls")
    output = Path("../private/results/1249A_aricles.xlsx")

    worker = TableWorker(None, input)
    worker.open_and_clean()
    if worker.tables is None or not bool(worker.tables):
        print("Errors with table")

    res = worker.simple_parser()
    if all(map(lambda x: x.size > 0, res)):
        print(res)

    unique_materials = set()
    for table in res:
        if table.empty:
            continue
        unique_materials |= set(table.material)
    unique_materials = list(unique_materials)

    # parsing materials (finding unique parts)
    unique_materials_dict = dict()
    unique_parts = set()
    parser = ParserV2(DELIMETERS)
    for material in unique_materials:
        parse_results = parser.parse(material)
        unique_materials_dict[material] = parse_results
        unique_parts |= set(parse_results.parts)

    wb = make_callculation_xlsx(res, unique_materials_dict)
    if wb is not None:
        wb.save(output)


def test_v2_parser():
    input = Path("../private/tables/1249A.xls")

    wb = xlrd.open_workbook(str(input), formatting_info=False)
    sheets = wb.sheet_names()
    sheet_tables = []
    text_sheets = []
    for sheetname in sheets:
        sheet = wb[sheetname]
        text_sheet = ""

        table = tp2.Table()
        for row in range(sheet.nrows):
            for col in range(sheet.ncols):
                val = sheet.cell_value(row, col)
                cell = tp2.Cell(row, col, val)
                if cell is None:
                    text_sheet += " "
                    continue
                text_sheet += cell.code
                table.set_cell(cell)
            text_sheet += "\n"
        sheet_tables.append(table)
        text_sheets.append(text_sheet)

    parser1 = re.compile(r"^\s*([RWHA])$")
    for sheet in text_sheets:
        lines = sheet.split("\n")
        print(lines)
        for line in lines:
            re.match(r"")


def main():
    # testfile = Path("../../private/tables/1108A.xls")
    inputs = Path("../private/tables")
    output = Path("../private/results/tables")

    parsed_cnt = 0
    all_cnt = 0
    error_files = []
    for file in inputs.iterdir():
        print("\n\n", file)
        all_cnt += 1

        worker = TableWorker(None, file)
        worker.open_and_clean()
        if worker.tables is None or not bool(worker.tables):
            print("Errors with table")
            error_files.append(file)
            continue

        res = worker.simple_parser()
        if all(map(lambda x: x.size > 0, res)):
            # print(res)
            parsed_cnt += 1

        unique_materials = set()
        for table in res:
            if table.empty:
                continue
            unique_materials |= set(table.material)
        unique_materials = list(unique_materials)

        # parsing materials (finding unique parts)
        unique_materials_dict = dict()
        unique_parts = set()
        parser = ParserV2(DELIMETERS)
        for material in unique_materials:
            parse_results = parser.parse(material)
            unique_materials_dict[material] = parse_results
            unique_parts |= set(parse_results.parts)

        wb = make_request_xlsx(res, unique_materials_dict)
        if wb is not None:
            wb.save(output / Path(file.name))
        wb2 = make_callculation_xlsx(res, unique_materials_dict)
        if wb2 is not None:
            wb2.save(output / (Path(file.stem + "_article" + file.suffix)))

    if all_cnt == 0:
        print("No files in the dir")
    else:
        print(f"Parsed: {parsed_cnt}/{all_cnt} = {parsed_cnt / all_cnt * 100:.1f}%")
        print(error_files)


if __name__ == "__main__":
    # test_callculation_table()
    # main()
    test_v2_parser()
